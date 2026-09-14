"""
Management command para carregar a planilha do acervo (template "Inserir
Material", v8+) no Nou-Rau. Mantém compatibilidade com a planilha v3.1 (abas BDU).

Uso:
    python manage.py migrate_spreadsheet /caminho/para/planilha.xlsx --sheet "Inserir Material"
    python manage.py migrate_spreadsheet /caminho/para/planilha.xlsx --sheet "Inserir Material" --dry-run --skip-red
    python manage.py migrate_spreadsheet /caminho/para/planilha.xlsx --sheet "Eventos"

Importador ESTRITO (taxonomia v12, set/2026) — cada linha é recusada (erro de
linha, savepoint desfeito, as demais seguem) quando:
  - a "Coleção" não casa uma coleção raiz (nunca cai na primeira raiz);
  - o "Tipo de informação" está fora do vocabulário canônico da coleção
    (Documentos Normativos e Vídeos passam a ser recusados; grafias legadas
    como "Acórdão"/"deliberacao" são normalizadas por taxonomy_v6.tipo_canonico);
  - Categoria, Subcategoria, Microcategoria ou Assunto preenchidos não resolvem.
Tipo novo NÃO é criado por padrão: `--allow-new-types` restaura o comportamento
antigo, só como exceção documentada. Aliases de grafia (contados no resumo):
"PLANO ANUAL DE CONTRATAÇÕES (PCA)" → "PLANO DE CONTRATAÇÕES ANUAL (PCA)" e
"FASE PREPARATÓRIA - X" → "X". Em `--dry-run`, TODAS as linhas recusadas são
listadas com motivo (insumo para a curadoria).
"""

import os
import re
import unicodedata
from collections import Counter

import openpyxl
import psycopg2
from django.core.management.base import BaseCommand, CommandError

from catalog import qualidade
from catalog.taxonomy_v6 import COLECOES_V6, tipo_canonico, tipos_de_colecao

# Abas de dados na planilha v3.1 (ordem de processamento)
DATA_SHEETS = ["Eventos", "Livros Digitais", "Trabalhos Acadêmicos", "Materiais Pedagógicos"]


class LinhaRecusadaError(ValueError):
    """Linha da planilha recusada pelo importador estrito (o motivo é a mensagem)."""


# Aliases de grafia — chaves e valores NORMALIZADOS (sem acento, minúsculas,
# espaços colapsados). Usados antes de casar com o banco; cada uso é contado
# em Command.alias_hits e sai no resumo final.
CATEGORIA_ALIASES = {
    # grafia das listas de apoio do template v8 (BDLP_Template_Insercao_v8 ...)
    "plano anual de contratacoes (pca)": "plano de contratacoes anual (pca)",
}
# Prefixo redundante retirado das subcategorias de Planejamento na v12
# (planilha "PARA CORREÇÃO" ainda traz "FASE PREPARATÓRIA - ETP" etc.).
SUBCATEGORIA_PREFIXOS_ALIAS = ("fase preparatoria - ", "fase preparatoria – ", "fase preparatoria-")

# Mapeamento de colunas da planilha v3.1 para campos internos.
# As chaves são os nomes EXATOS dos cabeçalhos da planilha.
COLUMN_MAP = {
    "Assunto": "assunto",
    "Categoria": "categoria",
    "Subcategoria": "subcategoria",
    "Microcategoria": "microcategoria",
    "Natureza": "natureza",
    "Coleção": "colecao",
    "Tipo de informação": "tipo_informacao",
    "Autor Principal": "author",
    "Título Principal": "title",
    "Título variante/outro idioma": "title_en",
    "Autoridade Intelectual": "autor_principal",
    "Assunto em português": "keywords",
    "Assunto em outro idioma": "keywords_en",
    "Resumo": "abstract",
    "Abstract": "abstract_en",
    "Nota": "description",
    "Edição": "edicao",
    "Apresentações do Evento": "event_description",
    "Imprenta": "source",
    "Descrição Física": "descricao_fisica",
    "ISSN/ISBN": "nlspi",
    "Identificador do Objeto Digital (DOI)": "doi",
    "Acesso Eletrônico": "acesso_eletronico",
    "Inserir uma Capa": "capa",
    "Permissão de acesso ao material": "tacesso_raw",
    "Tipologia": "tipologia",
    "Complexidade": "complexidade",
    "Aplicabilidade": "uso_futuro",
    "Uso Futuro": "uso_futuro",
    "Método": "metodo",
    "Resultados": "resultado",
    "Referências": "referencias",
    "Ano": "year",
}


def strip_accents(s):
    """Remove acentos de uma string para comparação fuzzy."""
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")


def normalize_text(val):
    """Limpa e normaliza texto da planilha."""
    if val is None:
        return ""
    return str(val).strip()


def _key(val):
    """Chave de comparação: sem acento, minúsculas, espaços colapsados."""
    return " ".join(strip_accents(normalize_text(val)).lower().split())


_COLECAO_POR_KEY = {_key(c["nome"]): c["nome"] for c in COLECOES_V6}


def extract_year(val):
    """Extrai ano de um valor."""
    val = normalize_text(val)
    match = re.search(r"\b(19|20)\d{2}\b", val)
    return match.group(0) if match else val


def generate_code(sequence_num):
    """Gera código único para o documento."""
    return f"bdlp-{sequence_num:06d}"


class Command(BaseCommand):
    help = "Migra planilha v3.1 de estudos (4 abas BDU) para o banco Nou-Rau"

    def add_arguments(self, parser):
        parser.add_argument("spreadsheet", type=str, help="Caminho para a planilha .xlsx")
        parser.add_argument("--dry-run", action="store_true", help="Apenas simula, sem inserir dados")
        parser.add_argument("--sheet", type=str, default=None, help="Nome de uma aba específica")
        parser.add_argument(
            "--skip-red", action="store_true",
            help="Pula linhas marcadas com fundo vermelho (curadoria sinaliza materiais a remover).",
        )
        parser.add_argument(
            "--start-seq", type=int, default=1,
            help="Sequência inicial dos códigos bdlp-XXXXXX (default 1). "
                 "Use para imports incrementais sem colidir com códigos existentes.",
        )
        parser.add_argument(
            "--allow-new-types", action="store_true",
            help="EXCEÇÃO documentada: cria em type_information um tipo fora do vocabulário "
                 "canônico em vez de recusar a linha (comportamento antigo). Não use em carga normal.",
        )

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.alias_hits = Counter()

    def reset_alias_hits(self):
        """Zera o contador de aliases usados (uma carga = uma contagem)."""
        self.alias_hits = Counter()

    def _get_write_connection(self):
        """Cria conexão de escrita usando o usuário php (não o portal_reader)."""
        return psycopg2.connect(
            dbname=os.environ.get("POSTGRES_DB", "nourau"),
            user=os.environ.get("POSTGRES_USER", "php"),
            password=os.environ["POSTGRES_PASSWORD"],  # sem fallback: falha claro se ausente
            host=os.environ.get("POSTGRES_HOST", "postgres"),
            port=os.environ.get("POSTGRES_PORT", "5432"),
        )

    def handle(self, *args, **options):
        spreadsheet_path = options["spreadsheet"]
        dry_run = options["dry_run"]
        sheet_name = options.get("sheet")
        skip_red = options.get("skip_red", False)
        start_seq = options.get("start_seq", 1)
        allow_new_types = options.get("allow_new_types", False)
        self.reset_alias_hits()
        if allow_new_types:
            self.stdout.write(self.style.WARNING(
                "--allow-new-types: tipos fora do vocabulário canônico serão CRIADOS (exceção documentada)."
            ))

        # read_only=False para ter acesso a fill (cor de fundo) — só carrega
        # com formatting quando precisamos detectar células vermelhas.
        try:
            wb = openpyxl.load_workbook(spreadsheet_path, read_only=not skip_red)
        except Exception as e:
            raise CommandError(f"Erro ao abrir planilha: {e}")

        # Determinar quais abas processar
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise CommandError(f"Aba '{sheet_name}' não encontrada. Disponíveis: {wb.sheetnames}")
            sheets_to_process = [sheet_name]
        else:
            sheets_to_process = [s for s in DATA_SHEETS if s in wb.sheetnames]
            if not sheets_to_process:
                raise CommandError(f"Nenhuma aba de dados encontrada. Disponíveis: {wb.sheetnames}")

        self.stdout.write(f"Abas a processar: {sheets_to_process}")

        # Detecção de linhas vermelhas (curadoria marca para remover/excluir)
        red_rows_by_sheet = {}
        if skip_red:
            self.stdout.write("Detectando linhas vermelhas (curadoria)...")
            for sn in sheets_to_process:
                red_rows_by_sheet[sn] = self._detect_red_rows(wb[sn])
                if red_rows_by_sheet[sn]:
                    self.stdout.write(f"  [{sn}] {len(red_rows_by_sheet[sn])} linhas vermelhas serão puladas")

        # Carregar mapeamentos do banco
        conn = self._get_write_connection()
        try:
            topic_map = self._load_topic_map(conn)
            category_map = self._load_category_map(conn)
            type_info_map = self._load_type_information_map(conn)
            assunto_map = self._load_assunto_map(conn)
            subcategoria_map = self._load_subcategoria_map(conn)
            microcategoria_map = self._load_microcategoria_map(conn)

            self.stdout.write(f"Coleções raiz: {sorted(nome for (parent, nome) in topic_map if parent == 0)}")
            self.stdout.write(f"Categorias: {len(category_map)}")
            self.stdout.write(f"Tipos de informação: {len(type_info_map)}")
            self.stdout.write(
                f"Taxonomia: {len(assunto_map)} assuntos, "
                f"{len(subcategoria_map)} subcategorias, "
                f"{len(microcategoria_map)} microcategorias"
            )

            global_seq = start_seq  # Contador global de código (continua import incremental)
            total_inserted = 0
            total_skipped = 0
            total_errors = []
            # Só no --dry-run: linhas (aceitas e recusadas) para as verificações de
            # qualidade (catalog.qualidade) — avisos para a curadoria, nunca recusa.
            registros_qualidade = []

            for sheet in sheets_to_process:
                ws = wb[sheet]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(row)

                if len(rows) < 2:
                    self.stdout.write(self.style.WARNING(f"\n[{sheet}] Aba vazia, pulando."))
                    continue

                # Detectar layout do template "Inserir Material":
                #   L1 = seções ('CLASSIFICAÇÃO' / 'METADADOS BDU')
                #   L2 = nomes das colunas (Assunto, Categoria, ...)
                #   L3 = indicadores '✱ OBRIGATÓRIO' (não é dado)
                #   L4+ = dados
                # Para BDU: L1 = headers, L2+ = dados (defaults).
                if sheet == "Inserir Material":
                    header_idx = 1   # 0-based: row 2 da planilha
                    data_start = 3   # 0-based: row 4 da planilha
                else:
                    header_idx = 0
                    data_start = 1

                # Mapear cabeçalhos
                header = [normalize_text(h) for h in rows[header_idx]]
                col_indices = self._map_headers(header)

                mapped = len(col_indices)
                total_cols = len(COLUMN_MAP)
                self.stdout.write(f"\n[{sheet}] Colunas mapeadas: {mapped}/{total_cols}")
                if missing := set(COLUMN_MAP.values()) - set(col_indices.keys()):
                    # Filtrar opcionais que não existem em todas as abas
                    optional = {"edicao", "event_description", "nlspi", "tipologia"}
                    real_missing = missing - optional
                    if real_missing:
                        self.stdout.write(self.style.WARNING(f"  Colunas não encontradas: {real_missing}"))

                # Filtrar linhas com dados (mantém row_num original p/ skip-red)
                red_rows = red_rows_by_sheet.get(sheet, set())
                data_rows = []
                for offset, row in enumerate(rows[data_start:], start=data_start):
                    excel_row = offset + 1  # 1-based: linha real na planilha
                    if not any(v is not None for v in row):
                        continue
                    if excel_row in red_rows:
                        continue  # curadoria marcou pra remover
                    data_rows.append((excel_row, row))

                red_skipped = len(red_rows) if red_rows else 0
                self.stdout.write(
                    f"  Registros com dados: {len(data_rows)}"
                    + (f" (+ {red_skipped} vermelhos pulados)" if red_skipped else "")
                )

                sheet_inserted = 0
                sheet_skipped = 0
                sheet_errors = []

                for row_num, row in data_rows:
                    record = self._parse_row(row, col_indices)
                    if not record.get("title"):
                        sheet_skipped += 1
                        continue
                    if dry_run:
                        registros_qualidade.append(qualidade.registro(
                            f"L{row_num}", record.get("title"), doi=record.get("doi"),
                            url=record.get("acesso_eletronico"), resumo=record.get("abstract"),
                            colecao=record.get("colecao"), tipo=record.get("tipo_informacao"),
                            assunto=record.get("assunto"), categoria=record.get("categoria"),
                            subcategoria=record.get("subcategoria"), autor=record.get("author"),
                        ))

                    # Savepoint por linha: um registro com erro não aborta a
                    # transação inteira (sem isto, um erro deixa a transação em
                    # estado abortado e TODAS as linhas seguintes falhariam com
                    # InFailedSqlTransaction).
                    with conn.cursor() as cur:
                        cur.execute("SAVEPOINT row_sp")
                    try:
                        topic_id = self._resolve_topic(record, topic_map)

                        # === Roteamento v8+ — campos diretos da planilha (sem de-para);
                        # cada resolução recusa a linha (LinhaRecusadaError) quando não casa.
                        type_info_id = self._ensure_type(
                            conn, record.get("tipo_informacao", ""), type_info_map,
                            allow_create=allow_new_types, dry_run=dry_run,
                        )
                        category_id = self._resolve_category(record.get("categoria", ""), category_map)
                        assunto_id = self._resolve_assunto(record.get("assunto", ""), assunto_map)
                        subcategoria_id = self._resolve_subcategoria(
                            record.get("subcategoria", ""), category_id, subcategoria_map
                        )
                        microcategoria_id = self._resolve_microcategoria(
                            record.get("microcategoria", ""), subcategoria_id, microcategoria_map
                        )

                        if dry_run:
                            title_preview = record["title"][:80]
                            self.stdout.write(f"  [DRY] #{global_seq} {sheet}/{row_num}: {title_preview}")
                        else:
                            code = generate_code(global_seq)
                            self._insert_document(
                                conn, record, code,
                                topic_id, category_id, type_info_id,
                                assunto_id, subcategoria_id, microcategoria_id,
                            )
                        sheet_inserted += 1
                        global_seq += 1
                        with conn.cursor() as cur:
                            cur.execute("RELEASE SAVEPOINT row_sp")

                    except Exception as e:
                        with conn.cursor() as cur:
                            cur.execute("ROLLBACK TO SAVEPOINT row_sp")
                            cur.execute("RELEASE SAVEPOINT row_sp")
                        titulo = (record.get("title") or "")[:60]
                        sheet_errors.append((sheet, row_num, f"{e} | {titulo}"))
                        # Em --dry-run, TODAS as recusas saem (insumo da curadoria); na
                        # carga real, as 5 primeiras aqui e as 20 primeiras no resumo.
                        if dry_run or len(sheet_errors) <= 5:
                            self.stdout.write(self.style.ERROR(f"  RECUSADA {sheet}/L{row_num}: {e} | {titulo}"))

                self.stdout.write(
                    f"  [{sheet}] Inseridos: {sheet_inserted}, "
                    f"Ignorados: {sheet_skipped}, Erros: {len(sheet_errors)}"
                )
                total_inserted += sheet_inserted
                total_skipped += sheet_skipped
                total_errors.extend(sheet_errors)

            if not dry_run:
                conn.commit()

        finally:
            conn.close()

        wb.close()

        self.stdout.write(self.style.SUCCESS("\n=== Resultado Final ==="))
        self.stdout.write(f"  Total inseridos: {total_inserted}")
        self.stdout.write(f"  Total ignorados: {total_skipped}")
        self.stdout.write(f"  Total recusados (erros de linha): {len(total_errors)}")
        self.stdout.write(
            "  Linhas que usaram alias de grafia: "
            f"categoria {self.alias_hits.get('categoria', 0)}, "
            f"subcategoria {self.alias_hits.get('subcategoria', 0)}"
        )

        if total_errors:
            limite = None if dry_run else 20
            rotulo = "Linhas recusadas (todas)" if dry_run else "Primeiras 20 linhas recusadas"
            self.stdout.write(self.style.ERROR(f"\n{rotulo}:"))
            for sheet, row_num, err in total_errors[:limite]:
                self.stdout.write(self.style.ERROR(f"  {sheet}/L{row_num}: {err}"))
            motivos = Counter(err.split(":", 1)[0] for _, _, err in total_errors)
            self.stdout.write("\n  Recusas por motivo:")
            for motivo, n in motivos.most_common():
                self.stdout.write(f"    {motivo}: {n}")

        if dry_run:
            achados = qualidade.analisar(registros_qualidade)
            self.stdout.write(self.style.WARNING(
                f"\nPossíveis redundâncias e problemas de qualidade — AVISOS para a curadoria, "
                f"não impedem a carga ({len(achados)} achados em {len(registros_qualidade)} linhas; "
                "semântica em catalog/qualidade.py e tools/db-refresh.md):"
            ))
            for linha in qualidade.formatar_relatorio(achados, limite=10):
                self.stdout.write(linha)
            self.stdout.write(self.style.WARNING("\n[DRY RUN] Nenhum dado foi inserido."))

    def _detect_red_rows(self, ws):
        """Retorna set com row_num (1-based) das linhas que contêm pelo menos
        uma célula com fundo vermelho (FF****00 ou similar).

        Usado para cumprir solicitação de curadoria: marcar materiais como
        vermelhos na planilha para que sejam excluídos do import.
        Verifica as primeiras 20 colunas (suficiente para Assunto, Título, etc.).
        """
        red_rows = set()
        max_check_col = min(ws.max_column, 20)
        for row_idx in range(2, ws.max_row + 1):  # pula header
            for col_idx in range(1, max_check_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                fill = cell.fill
                if fill is None or fill.fgColor is None:
                    continue
                rgb = fill.fgColor.rgb
                if not isinstance(rgb, str) or len(rgb) < 6:
                    continue
                rgb6 = rgb[-6:].upper()
                if not all(c in "0123456789ABCDEF" for c in rgb6):
                    continue
                r = int(rgb6[0:2], 16)
                g = int(rgb6[2:4], 16)
                b = int(rgb6[4:6], 16)
                # vermelho dominante: R alto, G/B baixos
                if r > 150 and g < 100 and b < 100:
                    red_rows.add(row_idx)
                    break
        return red_rows

    def _map_headers(self, header):
        """Mapeia cabeçalhos da planilha para campos internos com tolerância a acentos.

        Match EXATO tem prioridade sobre o substring — senão, na planilha v8,
        "Assunto" casaria com "Assunto em português (Palavras-chave)" e
        "Categoria" com "Subcategoria"/"Microcategoria".
        """
        col_indices = {}
        header_normalized = [strip_accents(h.lower()) if h else "" for h in header]

        for col_name, field_name in COLUMN_MAP.items():
            col_norm = strip_accents(col_name.lower())
            matched = False
            for i, h_norm in enumerate(header_normalized):
                if h_norm == col_norm:           # match exato — prioridade
                    col_indices[field_name] = i
                    matched = True
                    break
            if matched:
                continue
            for i, h_norm in enumerate(header_normalized):
                if h_norm and col_norm in h_norm:    # fallback substring
                    col_indices[field_name] = i
                    break
        return col_indices

    def _parse_row(self, row, col_indices):
        """Extrai e normaliza campos de uma linha da planilha."""
        record = {}
        for field, idx in col_indices.items():
            record[field] = normalize_text(row[idx]) if idx < len(row) else ""
        return record

    def _load_topic_map(self, conn):
        """Carrega os tópicos do banco como {(parent_id, nome normalizado): id}.

        A chave composta preserva subcoleções homônimas sob raízes diferentes
        (na janela da v12, "Enunciados" existe sob Jurisprudência e sob Doutrina).
        Raízes têm parent_id 0.
        """
        with conn.cursor() as cursor:
            cursor.execute("SELECT id, name, parent_id FROM topic")
            rows = cursor.fetchall()
        return {(row[2], _key(row[1])): row[0] for row in rows if row[1]}

    def _load_category_map(self, conn):
        """Carrega mapeamento nome→id de categorias do banco."""
        with conn.cursor() as cursor:
            cursor.execute("SELECT id, name FROM nr_category")
            rows = cursor.fetchall()
        return {row[1].strip().lower(): row[0] for row in rows}

    def _load_type_information_map(self, conn):
        """Carrega mapeamento nome→id de tipos de informação."""
        with conn.cursor() as cursor:
            cursor.execute("SELECT id, name FROM type_information")
            rows = cursor.fetchall()
        result = {}
        for row in rows:
            if row[1]:
                result[strip_accents(row[1].strip().lower())] = row[0]
        return result

    def _load_assunto_map(self, conn):
        """Mapeia nome (lower, sem acento)→id em nr_assunto."""
        with conn.cursor() as cursor:
            cursor.execute("SELECT id, nome FROM nr_assunto")
            rows = cursor.fetchall()
        return {strip_accents(r[1].strip().lower()): r[0] for r in rows if r[1]}

    def _load_subcategoria_map(self, conn):
        """Mapeia (category_id, nome_lower_sem_acento)→id em nr_subcategoria."""
        with conn.cursor() as cursor:
            cursor.execute("SELECT id, nome, category_id FROM nr_subcategoria")
            rows = cursor.fetchall()
        return {
            (r[2], strip_accents(r[1].strip().lower())): r[0]
            for r in rows if r[1]
        }

    def _load_microcategoria_map(self, conn):
        """Mapeia (subcategoria_id, nome_lower_sem_acento)→id em nr_microcategoria."""
        with conn.cursor() as cursor:
            cursor.execute("SELECT id, nome, subcategoria_id FROM nr_microcategoria")
            rows = cursor.fetchall()
        return {
            (r[2], strip_accents(r[1].strip().lower())): r[0]
            for r in rows if r[1]
        }

    def _resolve_colecao(self, record, topic_map):
        """(id da raiz, nome canônico da coleção) pela coluna 'Coleção'.

        Estrito: coleção vazia ou que não casa uma raiz do banco (e do vocabulário
        v12) recusa a linha — nunca cai na primeira raiz.
        """
        colecao_raw = normalize_text(record.get("colecao", ""))
        key = _key(colecao_raw)
        esperadas = ", ".join(c["nome"] for c in COLECOES_V6)
        if not key:
            raise LinhaRecusadaError(f"Coleção obrigatória vazia (esperado: {esperadas})")
        root_id = topic_map.get((0, key))
        colecao_nome = _COLECAO_POR_KEY.get(key)
        if root_id is None or colecao_nome is None:
            raise LinhaRecusadaError(f"Coleção não reconhecida: '{colecao_raw}' (esperado: {esperadas})")
        return root_id, colecao_nome

    def _tipo_da_colecao(self, record, colecao_nome):
        """Nome canônico v12 do 'Tipo de informação', validado contra a coleção.

        Aceita grafias legadas (taxonomy_v6.tipo_canonico); recusa a linha se o
        tipo estiver fora do vocabulário da coleção resolvida (Documentos
        Normativos e Vídeos inclusive) ou vazio.
        """
        tipo_raw = normalize_text(record.get("tipo_informacao", ""))
        aceitos = tipos_de_colecao(colecao_nome)
        canon = tipo_canonico(tipo_raw)
        if canon is None or canon not in aceitos:
            raise LinhaRecusadaError(
                f"Tipo de informação fora do vocabulário da coleção {colecao_nome}: "
                f"'{tipo_raw}' (aceitos: {', '.join(aceitos)})"
            )
        return canon

    def _resolve_topic(self, record, topic_map):
        """Resolve o topic_id: raiz pela 'Coleção' e subcoleção pelo 'Tipo de
        informação' canônico. Se o banco ainda não tiver a subcoleção (seção 1 do
        script de migração não aplicada), devolve a raiz."""
        root_id, colecao_nome = self._resolve_colecao(record, topic_map)
        canon = self._tipo_da_colecao(record, colecao_nome)
        sub_id = topic_map.get((root_id, _key(canon)))
        return sub_id if sub_id is not None else root_id

    def _ensure_type(self, conn, name, type_info_map, allow_create=False, dry_run=False):
        """Id do Tipo de Informação canônico v12 para o nome da planilha.

        - grafia legada → canônica (tipo_canonico); fora do vocabulário → recusa;
        - tipo canônico ausente do banco → recusa (falta a seção 1 do script);
        - `allow_create=True` (--allow-new-types) restaura a criação de tipo novo,
          como exceção documentada; em --dry-run só avisa (nada é escrito).
        """
        if not name or not name.strip():
            return None
        canon = tipo_canonico(name)
        if canon is None:
            if not allow_create:
                raise LinhaRecusadaError(
                    f"Tipo de informação fora do vocabulário canônico v12: '{name.strip()}' "
                    "(recusado; --allow-new-types só como exceção documentada)"
                )
            nome_final = name.strip()
        else:
            nome_final = canon
        key = _key(nome_final)
        if key in type_info_map:
            return type_info_map[key]
        if not allow_create:
            raise LinhaRecusadaError(
                f"Tipo de informação canônico ainda não existe no banco: '{nome_final}' "
                "(aplique a seção 1 de docker/postgres/migrations/2026-09-v12-taxonomia-e-busca.sql; "
                "--allow-new-types só como exceção documentada)"
            )
        if dry_run:
            self.stdout.write(self.style.WARNING(f"  [DRY] tipo novo seria criado: {nome_final}"))
            return None
        with conn.cursor() as cur:
            cur.execute("INSERT INTO type_information (name) VALUES (%s) RETURNING id", [nome_final])
            new_id = cur.fetchone()[0]
        type_info_map[key] = new_id
        return new_id

    def _resolve_category(self, categoria_name, category_map):
        """category_id pelo nome (alias de grafia → igualdade normalizada →
        substring >= 8 chars). Vazio → None; preenchido e não resolvido → recusa
        (nunca uma categoria-fallback que polua as facetas)."""
        if not categoria_name or not str(categoria_name).strip():
            return None
        key = _key(categoria_name)
        if key in CATEGORIA_ALIASES:
            key = CATEGORIA_ALIASES[key]
            self.alias_hits["categoria"] += 1
        for cat_key, cat_id in category_map.items():
            if _key(cat_key) == key:
                return cat_id
        # Busca parcial — só se substring é >= 8 chars para evitar
        # falsos positivos como "PCA" matching "PCA EM ANEXO".
        if len(key) >= 8:
            for cat_key, cat_id in category_map.items():
                cat_norm = _key(cat_key)
                if key in cat_norm or cat_norm in key:
                    return cat_id
        raise LinhaRecusadaError(f"Categoria não reconhecida: '{str(categoria_name).strip()}'")

    def _resolve_assunto(self, name, assunto_map):
        """assunto_id por nome (sem acento/caixa). Vazio → None; não resolvido → recusa."""
        if not name or not str(name).strip():
            return None
        key = _key(name)
        aid = assunto_map.get(key)
        if aid is None:
            raise LinhaRecusadaError(f"Assunto não reconhecido: '{str(name).strip()}'")
        return aid

    def _resolve_subcategoria(self, name, category_id, sub_map):
        """subcategoria_id sob category_id. Alias: tira o prefixo "FASE PREPARATÓRIA - ";
        depois igualdade normalizada e, por fim, substring (>= 5 chars) se casar
        UMA única subcategoria da categoria. Preenchido e não resolvido → recusa."""
        if not name or not str(name).strip() or not category_id:
            return None
        key = _key(name)
        for prefixo in SUBCATEGORIA_PREFIXOS_ALIAS:
            if key.startswith(prefixo):
                key = key[len(prefixo):].strip()
                self.alias_hits["subcategoria"] += 1
                break
        sub_id = sub_map.get((category_id, key))
        if sub_id:
            return sub_id
        if len(key) >= 5:
            candidatos = [
                sid for (cid, nome_norm), sid in sub_map.items()
                if cid == category_id and (key in nome_norm or nome_norm in key)
            ]
            if len(candidatos) == 1:
                return candidatos[0]
        raise LinhaRecusadaError(f"Subcategoria não reconhecida na categoria: '{str(name).strip()}'")

    def _resolve_microcategoria(self, name, subcategoria_id, mic_map):
        """microcategoria_id sob subcategoria_id (igualdade normalizada; substring
        >= 5 chars se casar uma única). Preenchido e não resolvido → recusa."""
        if not name or not str(name).strip() or not subcategoria_id:
            return None
        key = _key(name)
        mic_id = mic_map.get((subcategoria_id, key))
        if mic_id:
            return mic_id
        if len(key) >= 5:
            candidatos = [
                mid for (sid, nome_norm), mid in mic_map.items()
                if sid == subcategoria_id and (key in nome_norm or nome_norm in key)
            ]
            if len(candidatos) == 1:
                return candidatos[0]
        raise LinhaRecusadaError(f"Microcategoria não reconhecida na subcategoria: '{str(name).strip()}'")

    def _insert_document(
        self, conn, record, code,
        topic_id, category_id, type_info_id,
        assunto_id=None, subcategoria_id=None, microcategoria_id=None,
    ):
        """Insere um documento no banco via SQL direto."""
        # Montar description/info com campos auxiliares
        info_parts = []
        if record.get("nota"):
            info_parts.append(record["nota"])
        if record.get("tipo_informacao") and type_info_id is None:
            info_parts.append(f"Tipo de informação: {record['tipo_informacao']}")
        info = "\n".join(info_parts)

        # Autor principal: usar Autoridade Intelectual, fallback para Autor Principal
        autor_principal = record.get("autor_principal", "")
        if not autor_principal:
            autor_principal = record.get("author", "")

        # Etapa do processo licitatório: usar Subcategoria como rótulo livre (compatibilidade)
        etapa = record.get("subcategoria", "")[:255] if record.get("subcategoria") else ""

        # Ano: extrair como inteiro dedicado (preferir coluna `ano`)
        year_raw = record.get("year", "")
        year_str = extract_year(year_raw)
        try:
            ano_int = int(year_str) if year_str and year_str.isdigit() else None
        except (TypeError, ValueError):
            ano_int = None

        # Imprenta/source: se vazio, usar Ano
        source = record.get("source", "")
        if not source and year_str:
            source = year_str

        # Permissão: 'Aberto' (default), 'Restrito', ou texto cru se outro valor
        permissao_raw = (record.get("tacesso_raw") or "").strip()
        if permissao_raw.lower().startswith("aberto"):
            permissao = "Aberto"
        elif permissao_raw.lower().startswith("restrito"):
            permissao = "Restrito"
        else:
            permissao = permissao_raw[:20] if permissao_raw else None

        with conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO nr_document (
                    title, title_en, author, autor_principal,
                    abstract, abstract_en, keywords, keywords_en,
                    code, topic_id, category_id, status, remote,
                    acesso_eletronico, doi, nlspi, source,
                    tipologia, etapa_processo_licitatorio, complexidade,
                    uso_futuro, metodo, resultado, referencias,
                    description, info, owner_id,
                    typeinformation, typeinform_id,
                    descricao_fisica, edicao, event_description, capa,
                    assunto_id, subcategoria_id, microcategoria_id,
                    ano, permissao, natureza
                ) VALUES (
                    %s, %s, %s, %s,
                    %s, %s, %s, %s,
                    %s, %s, %s, 'a', 'y',
                    %s, %s, %s, %s,
                    %s, %s, %s,
                    %s, %s, %s, %s,
                    %s, %s, 1,
                    %s, %s,
                    %s, %s, %s, %s,
                    %s, %s, %s,
                    %s, %s, %s
                )
                """,
                [
                    record.get("title", "")[:1500],
                    record.get("title_en", "")[:1500],
                    record.get("author", "")[:3000],
                    autor_principal[:800],
                    record.get("abstract", ""),
                    record.get("abstract_en", ""),
                    record.get("keywords", ""),
                    record.get("keywords_en", ""),
                    code,
                    topic_id,
                    category_id,
                    record.get("acesso_eletronico", "")[:800],
                    record.get("doi", "")[:180],
                    record.get("nlspi", "")[:200],
                    source[:900],
                    record.get("tipologia", "")[:255],
                    etapa,
                    record.get("complexidade", "")[:50],
                    record.get("uso_futuro", ""),
                    record.get("metodo", ""),
                    record.get("resultado", ""),
                    record.get("referencias", ""),
                    "",  # description
                    info,
                    type_info_id,  # typeinformation (INT FK) — NULL if no match
                    type_info_id,  # typeinform_id (INT FK) — same
                    record.get("descricao_fisica", "")[:500],
                    record.get("edicao", "")[:900],
                    record.get("event_description", "")[:5000],
                    record.get("capa", "")[:65],
                    assunto_id,
                    subcategoria_id,
                    microcategoria_id,
                    ano_int,
                    permissao,
                    record.get("natureza", "")[:80],
                ],
            )
